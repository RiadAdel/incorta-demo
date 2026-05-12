import io

from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


LOGO_SVG = """
<svg id="Layer_1" data-name="Layer 1" xmlns="http://www.w3.org/2000/svg" viewBox="0 0 555.63 125.85" height="48" style="width:auto;display:block;">
  <defs><style>.cls-1{fill:#ffffff;}</style></defs>
  <path class="cls-1" d="M133.93,124.28h-25.43v-52.31c0-7.92-6.42-14.33-14.33-14.33s-14.33,6.42-14.33,14.33v52.31h-25.43v-52.31c0-21.96,17.8-39.76,39.76-39.76s39.76,17.8,39.76,39.76v52.31h0Z"/>
  <path class="cls-1" d="M233.19,91.7c-2.89,9.97-8.89,18.76-17.13,25.08-13.7,9.65-31.35,11.73-46.93,5.53-5.74-2.48-10.9-6.14-15.14-10.73-4.21-4.43-7.53-9.63-9.78-15.32-2.21-5.68-3.34-11.72-3.31-17.81-.02-6.05,1.17-12.04,3.5-17.62,2.38-5.6,5.8-10.69,10.08-15.01,4.27-4.44,9.38-7.99,15.04-10.44,5.63-2.35,11.68-3.54,17.77-3.48,10.28-.27,20.36,2.88,28.66,8.95,8.31,6.59,14.34,15.64,17.23,25.84h-27.67c-1.73-2.99-4.27-5.42-7.33-7.02-3.44-1.69-7.23-2.52-11.06-2.42-5.49-.1-10.77,2.08-14.6,6-3.83,4.12-5.9,9.58-5.78,15.2-.25,5.94,1.86,11.73,5.88,16.11,6.8,6.22,16.69,7.69,25,3.71,2.99-1.57,5.58-3.82,7.55-6.56h28.02Z"/>
  <path class="cls-1" d="M252.23,46.14c4.36-4.38,9.53-7.89,15.22-10.31,5.56-2.37,11.54-3.59,17.58-3.6,6.15-.05,12.25,1.07,17.97,3.29,11.57,4.65,20.85,13.66,25.82,25.1,2.45,5.63,3.73,11.69,3.75,17.83.08,6.3-1.06,12.55-3.36,18.41-2.29,5.65-5.72,10.78-10.07,15.06-4.36,4.43-9.55,7.96-15.28,10.38-5.87,2.4-12.17,3.61-18.51,3.55-6.34.09-12.63-1.11-18.49-3.52-5.74-2.49-10.89-6.14-15.14-10.73-4.21-4.35-7.54-9.48-9.78-15.11-2.22-5.64-3.34-11.65-3.31-17.71-.02-6.07,1.17-12.08,3.5-17.69,2.35-5.6,5.78-10.68,10.08-14.95ZM285.53,100.98c5.62-.01,10.99-2.33,14.87-6.4,3.99-4.16,6.2-9.71,6.17-15.47.04-5.81-2.18-11.4-6.17-15.62-7.75-8.21-20.68-8.59-28.9-.84-.29.27-.57.55-.84.84-4,4.21-6.21,9.81-6.17,15.62-.02,5.77,2.21,11.32,6.22,15.47,3.87,4.06,9.22,6.37,14.82,6.4h0Z"/>
  <path class="cls-1" d="M339.57,82.33v41.95h25.43v-41.95c.01-13.63,11.06-24.68,24.69-24.69v-25.43c-27.67.03-50.09,22.45-50.12,50.12Z"/>
  <path class="cls-1" d="M443.86,57.67v-25.48h-18.73V12.72h-25.43v111.54h25.43v-55.63c0-5.67,4.59-10.95,10.26-10.95h8.47Z"/>
  <path class="cls-1" d="M41.91,41.36v-9.15H0l.2,25.44h6.22c5.67,0,10.06,4.8,10.06,10.46v56.15h25.43V41.36h0Z"/>
  <circle class="cls-1" cx="29.71" cy="12.97" r="12.97"/>
  <path class="cls-1" d="M539.35,88.28v-56.2h-22.28v7.08c-2.36-1.49-4.86-2.75-7.45-3.78-5.74-2.23-11.85-3.35-18.01-3.3-6.05,0-12.05,1.22-17.62,3.59-5.7,2.43-10.89,5.93-15.26,10.32-4.32,4.28-7.75,9.37-10.1,14.98-2.33,5.61-3.52,11.64-3.51,17.72-.03,6.07,1.1,12.09,3.32,17.74,2.25,5.63,5.57,10.77,9.79,15.12,4.26,4.6,9.42,8.26,15.17,10.75,5.87,2.42,12.17,3.62,18.52,3.53,6.36.06,12.66-1.15,18.55-3.55,2.59-1.08,5.08-2.41,7.42-3.95v6.73h37.74v-26.42h-5.9c-5.73.02-10.39-4.62-10.39-10.35ZM507,94.55c-7.75,8.21-20.69,8.58-28.9.83-.28-.27-.56-.55-.83-.83-4.02-4.16-6.26-9.72-6.25-15.5-.04-5.82,2.18-11.42,6.19-15.64,7.75-8.23,20.71-8.61,28.93-.86.29.28.58.56.86.86,4.01,4.22,6.22,9.83,6.19,15.64.03,5.77-2.19,11.33-6.19,15.5h0Z"/>
</svg>
"""

INCORTA_SCALE = [
    [0.0,  "#3d1f8a"],
    [0.33, "#7c3aed"],
    [0.66, "#c084fc"],
    [1.0,  "#f59e0b"],
]

US_STATE_ABBREV = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR",
    "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE",
    "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS",
    "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS",
    "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV",
    "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY",
    "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK",
    "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT",
    "Vermont": "VT", "Virginia": "VA", "Washington": "WA", "West Virginia": "WV",
    "Wisconsin": "WI", "Wyoming": "WY",
}


def apply_chart_style(fig, categorical_y=False):
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font={"color": "#e8e6f0", "family": "sans-serif"},
        coloraxis_showscale=False,
        margin={"t": 10, "b": 40, "l": 10, "r": 10},
    )
    fig.update_xaxes(gridcolor="rgba(255,255,255,0.07)", zerolinecolor="rgba(255,255,255,0.1)")
    fig.update_yaxes(gridcolor="rgba(255,255,255,0.07)", zerolinecolor="rgba(255,255,255,0.1)")
    if categorical_y:
        fig.update_yaxes(categoryorder="total ascending")
    return fig


def build_pdf(df, customers, total_sales, avg_sales, avg_orders):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_fill_color(36, 28, 85)
    pdf.rect(0, 0, 210, 36, "F")
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(255, 255, 255)
    pdf.set_y(6)
    pdf.cell(0, 12, "Customer Sales Dashboard", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(196, 181, 253)
    pdf.cell(0, 10, "Online Store   |   Performance Analytics", align="C", new_x="LMARGIN", new_y="NEXT")

    pdf.set_y(45)

    kpi_labels = ["Total Customers", "Total Sales", "Avg Sales / Customer", "Avg Orders / Customer"]
    kpi_values = [f"{customers:,}", f"${total_sales:,.0f}", f"${avg_sales:,.0f}", f"{avg_orders:,.1f}"]
    col_w, x_start, y_kpi = 43, 15, pdf.get_y()
    for i, (label, value) in enumerate(zip(kpi_labels, kpi_values)):
        x = x_start + i * (col_w + 2)
        pdf.set_fill_color(240, 238, 255)
        pdf.set_draw_color(124, 58, 237)
        pdf.set_line_width(0.5)
        pdf.rect(x, y_kpi, col_w, 22, "FD")
        pdf.set_font("Helvetica", "", 7)
        pdf.set_text_color(124, 58, 237)
        pdf.set_xy(x + 2, y_kpi + 3)
        pdf.cell(col_w - 4, 5, label.upper())
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(36, 28, 85)
        pdf.set_xy(x + 2, y_kpi + 10)
        pdf.cell(col_w - 4, 8, value)

    pdf.set_y(y_kpi + 30)

    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(36, 28, 85)
    pdf.cell(0, 8, "Top Customers", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    headers = ["Account", "City", "State", "Orders", "Avg Order ($)", "Total Sales ($)"]
    widths  = [28, 32, 42, 18, 33, 35]
    aligns  = ["L", "L", "L", "R", "R", "R"]

    pdf.set_fill_color(36, 28, 85)
    pdf.set_text_color(255, 255, 255)
    pdf.set_draw_color(100, 80, 180)
    pdf.set_font("Helvetica", "B", 8)
    for h, w in zip(headers, widths):
        pdf.cell(w, 8, h, border=1, fill=True, align="C")
    pdf.ln()

    top_df = df.sort_values("Total_Sales", ascending=False).head(25)
    pdf.set_font("Helvetica", "", 7.5)
    for i, (_, row) in enumerate(top_df.iterrows()):
        if i % 2 == 0:
            pdf.set_fill_color(245, 243, 255)
        else:
            pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(50, 40, 100)
        vals = [
            str(row["AccountNumber"])[:14],
            str(row["City"])[:14],
            str(row["StateProvinceName"])[:18],
            str(int(row["Sales_Order_Count"])),
            f"${row['Average_Sales_Order']:,.0f}",
            f"${row['Total_Sales']:,.0f}",
        ]
        for val, w, align in zip(vals, widths, aligns):
            pdf.cell(w, 7, val, border=1, fill=True, align=align)
        pdf.ln()

    return bytes(pdf.output())


def build_pptx(df, customers, total_sales, avg_sales, avg_orders):
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def dark_bg(slide, rgb=(13, 10, 46)):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*rgb)

    def add_label(slide, text, x, y, w, h, size=14, bold=False,
                  color=(255, 255, 255), align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)

    s1 = prs.slides.add_slide(blank)
    dark_bg(s1, (36, 28, 85))
    add_label(s1, "Customer Sales Dashboard",
              Inches(1), Inches(2.5), Inches(11.33), Inches(1.5),
              size=40, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
    add_label(s1, "Online Store  ·  Performance Analytics",
              Inches(1), Inches(4.1), Inches(11.33), Inches(0.8),
              size=18, color=(196, 181, 253), align=PP_ALIGN.CENTER)

    s2 = prs.slides.add_slide(blank)
    dark_bg(s2)
    add_label(s2, "Key Performance Indicators",
              Inches(0.5), Inches(0.3), Inches(12), Inches(0.7),
              size=24, bold=True, color=(196, 181, 253))

    kpi_data = [
        ("Total Customers",      f"{customers:,}"),
        ("Total Sales",          f"${total_sales:,.0f}"),
        ("Avg Sales / Customer", f"${avg_sales:,.0f}"),
        ("Avg Orders / Customer",f"{avg_orders:,.1f}"),
    ]
    for i, (label, value) in enumerate(kpi_data):
        box = s2.shapes.add_textbox(Inches(0.3 + i * 3.2), Inches(1.4), Inches(3.0), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(36, 28, 85)
        box.line.color.rgb = RGBColor(124, 58, 237)
        box.line.width = Pt(1.5)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Pt(16)
        tf.margin_left = Pt(8)

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = label.upper()
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(196, 181, 253)

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(10)
        r2 = p2.add_run()
        r2.text = value
        r2.font.size = Pt(28)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(255, 255, 255)

    s3 = prs.slides.add_slide(blank)
    dark_bg(s3)
    add_label(s3, "Top Customers",
              Inches(0.5), Inches(0.3), Inches(12), Inches(0.7),
              size=24, bold=True, color=(196, 181, 253))

    top_df = df.sort_values("Total_Sales", ascending=False).head(15)
    col_hdrs = ["Account", "City", "State", "Orders", "Total Sales"]
    tbl = s3.shapes.add_table(
        len(top_df) + 1, len(col_hdrs),
        Inches(0.3), Inches(1.1), Inches(12.7), Inches(5.8)
    ).table

    for j, h in enumerate(col_hdrs):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(36, 28, 85)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.size = Pt(10)

    row_aligns = [PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT]
    for i, (_, row) in enumerate(top_df.iterrows()):
        vals = [
            str(row["AccountNumber"]),
            str(row["City"]),
            str(row["StateProvinceName"]),
            str(int(row["Sales_Order_Count"])),
            f"${row['Total_Sales']:,.0f}",
        ]
        bg = RGBColor(26, 19, 80) if i % 2 == 0 else RGBColor(18, 13, 55)
        for j, (v, align) in enumerate(zip(vals, row_aligns)):
            cell = tbl.cell(i + 1, j)
            cell.text = v
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.alignment = align
            if p.runs:
                p.runs[0].font.color.rgb = RGBColor(232, 230, 240)
                p.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
